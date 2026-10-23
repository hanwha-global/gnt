"""사용자 리서치 엑셀(공시 기반, FY2024)을 대시보드 기준 데이터(JSON)로 변환한다.

한 번만 돌리는 변환 스크립트이며, 이후 최신 수치는 update_dart.py 가 덮어쓴다.
사용: python scripts/build_seed.py "<국내증권사 리서치 ..._설립취득.xlsx>"
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gnd import classify_entity, normalize_country  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "site" / "data" / "domestic"

FIRMS = {
    "1) 하나증권": ("하나증권", "은행계"),
    "2) 신한증권": ("신한투자증권", "은행계"),
    "3) KB증권": ("KB증권", "은행계"),
    "4) NH투자증권": ("NH투자증권", "은행계"),
    "5) 미래에셋증권": ("미래에셋증권", "비은행계"),
    "6) 한국투자증권": ("한국투자증권", "비은행계"),
    "7) 삼성증권": ("삼성증권", "비은행계"),
    "8) 키움증권": ("키움증권", "비은행계"),
}

# 원자료 오기 보정: (법인명 일부, 필드, 보정값, 사유)
CORRECTIONS = [
    ("Korea Investment & Securities America", "country", "미국", "법인명상 미국 법인(원자료 소재지 '인도네시아' 오기)"),
    ("NH Investment & Securities(H.K.)", "country", "홍콩", "법인명상 홍콩 법인(원자료 '중국')"),
]

# 지분율·관계 메모(원자료 비고 기준)
RELATIONS = {
    "BIDV Securities": "관계사",
    "Wedbush Financial Services": "지분투자",
    "북경랑자하나자산관리": "지분투자",
}


def num(v):
    return round(float(v), 1) if isinstance(v, (int, float)) else None


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["진출방식"]
    entities, corrections, firm = [], [], None
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if not cells:
            continue
        if isinstance(cells[0], str) and cells[0].strip() in FIRMS:
            firm = FIRMS[cells[0].strip()]
            continue
        if firm is None or cells[0] == "해외법인명" or len(cells) < 6:
            continue
        name, parent, biz, country, est, mode = (str(c).strip() for c in cells[:6])
        rest = cells[6:]
        note = rest.pop(0) if rest and isinstance(rest[0], str) else ""
        fin = [num(v) for v in rest[:5]] + [None] * (5 - len(rest[:5]))
        status = "매각" if any(isinstance(v, str) and "매각" in v for v in rest) else "운영"
        e = {
            "firm": firm[0], "group": firm[1], "name": name, "parent": parent,
            "business": biz, "country": normalize_country(country),
            "established": est, "entry_mode": mode, "note": note,
            "assets": fin[0], "liabilities": fin[1], "equity": fin[2],
            "op_income": fin[3], "net_income": fin[4],
            "relation": next((v for k, v in RELATIONS.items() if k in name), "종속(자회사)"),
            "status": status, "basis": "FY2024 사업보고서",
        }
        for key, field, value, why in CORRECTIONS:
            if key in name and e[field] != value:
                corrections.append({"name": name, "field": field, "from": e[field], "to": value, "reason": why})
                e[field] = value
        e["kind"], e["exclude_reason"] = classify_entity(name, biz)
        entities.append(e)

    # 원자료 요약표(국가×증권사)의 사무소 집계와 토스증권 요약치
    offices = [
        {"firm": "KB증권", "country": "인도네시아", "count": 1},
        {"firm": "미래에셋증권", "country": "중국", "count": 2},
        {"firm": "한국투자증권", "country": "중국", "count": 2},
        {"firm": "삼성증권", "country": "중국", "count": 1},
        {"firm": "NH투자증권", "country": "중국", "count": 1},
        {"firm": "KB증권", "country": "중국", "count": 1},
        {"firm": "미래에셋증권", "country": "베트남", "count": 1},
        {"firm": "삼성증권", "country": "일본", "count": 1},
    ]
    for c in ("미국", "싱가포르", "일본"):
        entities.append({
            "firm": "토스증권", "group": "비은행계", "name": None, "parent": "토스증권",
            "business": None, "country": c, "established": None, "entry_mode": None,
            "note": "원자료 요약표에만 집계(법인 상세 미기재). DART 갱신 시 반영",
            "assets": None, "liabilities": None, "equity": None, "op_income": None, "net_income": None,
            "relation": "종속(자회사)", "status": "운영", "basis": "FY2024 요약표", "kind": "법인",
            "exclude_reason": None,
        })

    # 한화투자증권(비교 기준선): 법인명·국가만 두고 재무는 DART 수집값으로 채운다
    for name, country, biz, mode in (
        ("Pinetree Securities Corporation", "베트남", "증권업", "인수"),
        ("Pinetree Securities Pte. Ltd.", "싱가포르", "증권업", None),
        ("PT Ciptadana Sekuritas Asia", "인도네시아", "증권업", "인수"),
        ("PT Ciptadana Asset Management", "인도네시아", "자산운용업", "인수"),
    ):
        entities.append({
            "firm": "한화투자증권", "group": "비은행계", "name": name, "parent": "한화투자증권",
            "business": biz, "country": country, "established": None, "entry_mode": mode,
            "note": "", "assets": None, "liabilities": None, "equity": None, "op_income": None,
            "net_income": None, "relation": "종속(자회사)", "status": "운영",
            "basis": "DART 수집 대기", "kind": "법인", "exclude_reason": None,
        })

    s = wb["Summary"]
    firms = []
    for r in range(4, 14):
        name = s.cell(r, 3).value
        if isinstance(name, str) and name.endswith("증권"):
            firms.append({
                "firm": {"신한증권": "신한투자증권", "NH증권": "NH투자증권"}.get(name, name),
                "established": s.cell(r, 4).value.strftime("%Y-%m-%d"),
                "assets": s.cell(r, 5).value, "liabilities": s.cell(r, 6).value,
                "equity": s.cell(r, 7).value, "op_income": s.cell(r, 8).value,
                "net_income": s.cell(r, 9).value, "basis": "FY2024 사업보고서(연결)",
            })
    firms.append({"firm": "한화투자증권", "established": None, "assets": None, "liabilities": None,
                  "equity": None, "op_income": None, "net_income": None, "basis": "DART 수집 대기"})
    group = {f[0]: f[1] for f in FIRMS.values()}
    for f in firms:
        f["group"] = group.get(f["firm"], "비은행계")

    OUT.mkdir(parents=True, exist_ok=True)
    meta = {"source": "국내 주요 증권사 사업보고서(FY2024) 기반 사용자 리서치, 단위 백만원",
            "corrections": corrections}
    (OUT / "seed.json").write_text(json.dumps(
        {"meta": meta, "firms": firms, "entities": entities, "offices": offices},
        ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"entities={len(entities)} firms={len(firms)} corrections={len(corrections)}")


if __name__ == "__main__":
    main(sys.argv[1])
