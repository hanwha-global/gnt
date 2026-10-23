"""DART 공시(dart_subs.json) + 기준 정보(entities.json) + 출자 이력(investments.json) → 다운로드 파일.

    python hanwha/scripts/build_downloads.py           # 엑셀·1장 보고서(Word, LibreOffice가 있으면 PDF)를 만든다
    python hanwha/scripts/build_downloads.py --check   # 다시 만들어야 하면 rebuild=true, 아니면 rebuild=false 출력

만드는 파일: downloads/해외자회사_공시실적.xlsx, downloads/해외자회사_공시실적_1장보고서.docx(.pdf).
파일 이름에 기간을 넣지 않아 화면 링크가 바뀌지 않는다. 기간은 제목과 기준 줄에 적힌다.
공시 기준(보고서 접수번호·공시 숫자·다음 공시 항목·출자내역·기준 정보·이 스크립트)이 바뀐 날만 다시 만들도록
downloads/basis.json에 기준을 남긴다. PDF까지 만든 경우에만 기준을 남겨, 실패하면 다음 실행에서 다시 만든다.
공시 숫자만 쓰며 내부 계획·현지 결산·비공시 특이사항은 넣지 않는다. LLM을 쓰지 않는다.
"""
import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA, DL = ROOT / "data", ROOT / "downloads"
XLSX = DL / "해외자회사_공시실적.xlsx"
DOCX = DL / "해외자회사_공시실적_1장보고서.docx"
PDF = DOCX.with_suffix(".pdf")
BASIS = DL / "basis.json"
NAVY, ORANGE, GREY = "#1F2A44", "#F37321", "#666666"
FONT = "맑은 고딕"


def load(name):
    p = DATA / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()[:16] if p.exists() else None


def next_filing(reg):
    """아직 완료되지 않은 가장 이른 정기보고서 항목(1장 보고서 '다음 공시' 줄)."""
    xs = sorted((t for t in (reg or {}).get("todo", []) if t["id"].startswith("periodic-") and t["status"] != "done" and t.get("due")),
                key=lambda t: t["due"])
    return xs[0] if xs else None


def basis():
    d = load("dart_subs.json") or {}
    P = d.get("periods", [])
    used = {k: d.get(k) for k in ("periods", "stakes", "invest")}  # 같은 접수번호라도 숫자가 바뀌면 다시 만든다
    return {"periods": [p["report"]["rcept_no"] for p in P], "label": P[-1]["label"] if P else None,
            "invest": (d.get("invest") or {}).get("report", {}).get("rcept_no"),
            "data": hashlib.sha256(json.dumps(used, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16],
            "next": (next_filing(load("regulatory.json")) or {}).get("id"),
            "investments": sha(DATA / "investments.json"), "entities": sha(DATA / "entities.json"),
            "script": sha(Path(__file__))}


def needs_rebuild():
    if not all(p.exists() for p in (XLSX, DOCX, PDF, BASIS)):
        return True
    return json.loads(BASIS.read_text(encoding="utf-8")) != basis()


# ---------- 화면(app.js model)과 같은 계산 ----------
def model():
    D, ents, inv, reg = load("dart_subs.json"), load("entities.json")["entities"], load("investments.json"), load("regulatory.json")
    P = D["periods"]
    L = P[-1]
    F = next((p for p in reversed(P) if p["period"] == "FY" and p is not L), None)  # 최신이 사업보고서뿐이면 FY 열 없음

    def row(p, k):
        x = (p or {}).get("subs", {}).get(k)
        if not x:
            return None
        eq = x["assets"] - x["liab"] if x.get("assets") is not None and x.get("liab") is not None else None
        roe = x["ni"] * (12 / p["months"]) / eq if x.get("ni") is not None and eq else None
        return {**x, "eq": eq, "roe": roe}

    E = [{**e, "stake": D.get("stakes", {}).get(e["key"], 100) / 100, "stake_ok": e["key"] in D.get("stakes", {}),
          "l": row(L, e["key"]), "f": row(F, e["key"]),
          "inv": ((D.get("invest") or {}).get("rows") or {}).get(e["key"]),
          "hist": [x for x in (inv or {}).get("rows", []) if x["key"] == e["key"]]} for e in ents]

    def tot(side, m, w=False):
        xs = [(e[side][m], e["stake"] if w else 1) for e in E if e[side] and e[side].get(m) is not None]
        return sum(v * s for v, s in xs) if xs else None

    def roe(side, p):
        xs = [e[side] for e in E if e[side] and e[side].get("ni") is not None and e[side]["eq"]]
        eq = sum(x["eq"] for x in xs)
        return sum(x["ni"] for x in xs) * (12 / p["months"]) / eq if eq else None

    def both(fa, fb):
        """증감·비율 합계는 두 값이 모두 있는 법인끼리만 더한다(한쪽에만 있는 법인이 섞이지 않게)."""
        xs = [(fa(e), fb(e)) for e in E if fa(e) is not None and fb(e) is not None]
        return (sum(a for a, _ in xs), sum(b for _, b in xs)) if xs else None

    return {"D": D, "L": L, "F": F, "E": E, "tot": tot, "roe": roe, "both": both, "inv": inv, "next": next_filing(reg)}


def eq_of(e, side):
    return (e[side] or {}).get("eq")


def nav_of(e):
    return eq_of(e, "l") * e["stake"] if eq_of(e, "l") is not None else None


def book_of(e):
    return (e["inv"] or {}).get("end_book")


def vsum(vals):
    vals = [v for v in vals if v is not None]
    return sum(vals) if vals else None


def fmt(v):
    return "–" if v is None else (f"{v:,.1f}" if v >= 0 else f"({abs(v):,.1f})")


def pct(v, d=1):
    return "–" if v is None else f"{v * 100:.{d}f}%"


def ymd(s):
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if s else ""


def summary_lines(M):
    """1장 보고서·엑셀 첫 줄에 쓰는 문장. 숫자에서 바로 만든다(판단·전망 문구 없음)."""
    L, F, E, tot, roe = M["L"], M["F"], M["E"], M["tot"], M["roe"]
    a = (f"{L['label']} 영업수익 {fmt(tot('l', 'rev'))} · 순이익 {fmt(tot('l', 'ni'))} · 지배지분 순이익 {fmt(tot('l', 'ni', True))}"
         f" · {'' if L['months'] == 12 else '연환산 '}ROE {pct(roe('l', L))}")
    if F:
        a += f"   │   {F['label']} 영업수익 {fmt(tot('f', 'rev'))} · 지배지분 순이익 {fmt(tot('f', 'ni', True))} · ROE {pct(roe('f', F))}"
    prof = [e for e in E if e["l"] and (e["l"]["ni"] or 0) > 0]
    loss = sorted((e for e in E if e["l"] and (e["l"]["ni"] or 0) < 0), key=lambda e: e["l"]["ni"])
    b = f"{L['label']} 흑자 {len(prof)}곳" + (f"({', '.join(e['label'] + ' ' + fmt(e['l']['ni']) for e in prof)})" if prof else "")
    if loss:
        b += f", 손실 {len(loss)}곳({', '.join(e['label'] + ' ' + fmt(e['l']['ni']) for e in loss)})."
    acq = [e["label"] for e in E if e["f"] and e["f"].get("ni") is None]
    gone = [e["label"] for e in E if not e["f"]]
    if F and acq:
        b += f" {', '.join(acq)}은 {F['label']} 말 취득이라 사업보고서 손익이 없음."
    if F and gone:
        b += f" {', '.join(gone)}은 {F['label']} 요약재무정보에 없음."
    nost = [e["label"] for e in E if not e["stake_ok"]]
    if nost:
        b += f" {', '.join(nost)}은 공시 지분율이 없어 100%로 계산."
    return a, b


# ---------- 엑셀 ----------
def build_xlsx(M):
    import xlsxwriter

    L, F, E, tot, roe, D = M["L"], M["F"], M["E"], M["tot"], M["roe"], M["D"]
    wb = xlsxwriter.Workbook(str(XLSX))
    fm = lambda **k: wb.add_format({"font_name": FONT, "font_size": 10, "valign": "vcenter", **k})  # noqa: E731
    title, sub = fm(bold=True, font_size=15, font_color=NAVY), fm(font_color=GREY, font_size=9)
    hdr = fm(bold=True, bg_color=NAVY, font_color="white", align="center", border=1, text_wrap=True)
    lab, labb = fm(border=1, indent=1), fm(border=1, bold=True, indent=1)
    num = fm(border=1, num_format="#,##0.0;[Red](#,##0.0);-")
    numb = fm(border=1, bold=True, num_format="#,##0.0;[Red](#,##0.0);-", bg_color="#F2F2F2")
    qty = fm(border=1, num_format="#,##0;[Red](#,##0);-")
    pc = fm(border=1, num_format="0.0%;[Red]-0.0%;-", align="right")
    st = fm(border=1, num_format="0.00", align="right")
    pcb = fm(border=1, bold=True, num_format="0.0%;[Red]-0.0%;-", align="right", bg_color="#F2F2F2")
    txt = fm(border=1, align="center")
    note = fm(font_size=9, font_color=GREY)
    link = fm(font_color="#1F4E9E", underline=1)

    dash = fm(border=1, align="right", font_color=GREY)

    def put(ws, r, c, v, f):
        if v is None:
            ws.write_string(r, c, "–", dash)
        elif isinstance(v, str):
            ws.write_string(r, c, v, f)
        else:
            ws.write_number(r, c, v, f)

    def sheet(name, ttl, basis_line, widths):
        ws = wb.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.set_column(0, 0, 2)
        for i, w in enumerate(widths):
            ws.set_column(i + 1, i + 1, w)
        ws.write(0, 1, ttl, title)
        ws.write(1, 1, basis_line, sub)
        ws.set_landscape(); ws.set_paper(9); ws.fit_to_pages(1, 0)
        return ws

    src = f"{F['report']['name']} · {L['report']['name']}" if F else L["report"]["name"]
    # ① 실적 요약
    ws = sheet("①실적 요약", f"해외자회사 공시 실적 – {L['label']}", f"기준 DART {src} · 단위: KRW 백만원, ROE % · 종속기업 요약재무정보(법인 단위, 연결 조정 전)",
               [26, 9] + [12] * 7)
    heads = ["법인", "지분율"] + ([f"{F['label']}\n영업수익", f"{F['label']}\n순이익", f"{F['label']}\nROE"] if F else []) + \
            [f"{L['label']}\n영업수익", f"{L['label']}\n순이익", f"{L['label']}\n지배지분 순이익", f"{L['label']}\nROE" + ("" if L["months"] == 12 else "(연환산)")]
    ws.set_row(3, 32)
    for c, h in enumerate(heads):
        ws.write(3, 1 + c, h, hdr)
    r = 4
    for e in E:
        f, l = e["f"] or {}, e["l"] or {}
        vals = [e["stake"]] + ([f.get("rev"), f.get("ni"), f.get("roe")] if F else []) + \
               [l.get("rev"), l.get("ni"), None if l.get("ni") is None else round(l["ni"] * e["stake"], 1), l.get("roe")]
        ws.write(r, 1, e["label"], labb)
        for c, v in enumerate(vals):
            is_pct = c == 0 or (F and c == 3) or c == len(vals) - 1
            put(ws, r, 2 + c, v, fm(border=1, num_format="0.00%", align="center") if c == 0 else pc if is_pct else num)
        r += 1
    ws.write(r, 1, "4사 합산 (100%)", labb)
    ws.write_blank(r, 2, None, numb)
    tv = ([tot("f", "rev"), tot("f", "ni"), roe("f", F)] if F else []) + [tot("l", "rev"), tot("l", "ni"), tot("l", "ni", True), roe("l", L)]
    for c, v in enumerate(tv):
        put(ws, r, 3 + c, v, pcb if (F and c == 2) or c == len(tv) - 1 else numb)
    a, b = summary_lines(M)
    ws.write(r + 2, 1, a, note)
    ws.write(r + 3, 1, b, note)
    ws.write(r + 4, 1, f"ROE = 누적 순이익 × (12 ÷ 개월 수) ÷ 기말 자본(자산 − 부채). {L['label']}은 {L['months']}개월 누적. 지배지분 순이익 = 순이익 × 지분율(타법인 출자현황).", note)
    ws.freeze_panes(4, 2)

    # ② 재무상태
    per = ([F] if F else []) + [L]
    ws = sheet("②재무상태", "재무상태 (DART 공시)", f"기준 {' · '.join(p['label'] + ' 말' for p in per)} · 단위: KRW 백만원", [26] + [12] * (3 * len(per) + 1))
    ws.set_row(3, 30)
    ws.write(3, 1, "법인", hdr)
    c = 2
    for p in per:
        for m in ("자산", "부채", "자본"):
            ws.write(3, c, f"{p['label']} 말\n{m}", hdr); c += 1
    if F:
        ws.write(3, c, "자본 증감", hdr)
    for i, e in enumerate(E):
        ws.write(4 + i, 1, e["label"], labb)
        c = 2
        for side in (["f"] if F else []) + ["l"]:
            for m in ("assets", "liab", "eq"):
                put(ws, 4 + i, c, (e[side] or {}).get(m), num); c += 1
        if F:
            put(ws, 4 + i, c, eq_of(e, "l") - eq_of(e, "f") if eq_of(e, "l") is not None and eq_of(e, "f") is not None else None, num)
    rt = 4 + len(E)
    ws.write(rt, 1, "4사 합산", labb)
    c = 2
    for side in (["f"] if F else []) + ["l"]:
        for m in ("assets", "liab", "eq"):
            put(ws, rt, c, tot(side, m), numb); c += 1
    if F:
        dT = M["both"](lambda e: eq_of(e, "l"), lambda e: eq_of(e, "f"))
        put(ws, rt, c, dT[0] - dT[1] if dT else None, numb)

    # ③ 출자 현황 (DART 타법인 출자현황)
    ir = D.get("invest")
    if ir:
        ws = sheet("③출자 현황(DART)", "출자 현황 (DART 타법인 출자현황)", f"기준 {ir['report']['name']} (접수 {ymd(ir['report']['rcept_dt'])}) · 단위: KRW 백만원, 주, %",
                   [26, 10, 11, 11, 13, 8, 11, 12, 11, 10, 13, 8, 11, 12, 12])
        top = [("법인", 1), ("출자목적", 1), ("최초 취득", 2), ("기초 잔액", 3), ("증가(감소)", 3), ("기말 잔액", 3), ("최근 사업연도", 2)]
        c = 1
        for t, n in top:
            if n == 1:
                ws.merge_range(3, c, 4, c, t, hdr)
            else:
                ws.merge_range(3, c, 3, c + n - 1, t, hdr)
            c += n
        for c, t in enumerate(["일자", "금액", "수량", "지분율", "장부가", "수량", "취득(처분)", "평가손익", "수량", "지분율", "장부가", "총자산", "당기순손익"]):
            ws.write(4, 3 + c, t, hdr)
        for i, e in enumerate(E):
            x = e["inv"] or {}
            rr = 5 + i
            ws.write(rr, 1, e["label"], labb)
            put(ws, rr, 2, x.get("purpose"), txt)
            put(ws, rr, 3, (x.get("first_date") or "").replace(".", "-") or None, txt)
            for c, (k, f) in enumerate([("first_amount", num), ("begin_qty", qty), ("begin_stake", st), ("begin_book", num), ("chg_qty", qty),
                                        ("chg_amount", num), ("chg_valuation", num), ("end_qty", qty), ("end_stake", st), ("end_book", numb),
                                        ("recent_assets", num), ("recent_ni", num)]):
                put(ws, rr, 4 + c, x.get(k), f)
        rt = 5 + len(E)
        ws.write(rt, 1, "4사 합계 (장부가)", labb)
        put(ws, rt, 7, vsum((e["inv"] or {}).get("begin_book") for e in E), numb)
        put(ws, rt, 13, vsum(book_of(e) for e in E), numb)
        ws.write(rt + 2, 1, "DART 원문 표 구성을 그대로 따름. 장부가는 당사 별도 재무제표 기준. '최근 사업연도'는 직전 사업연도 말 피투자회사 재무.", note)

    # ④ 출자 이력 (내부 출자내역, 수동)
    inv = M["inv"]
    if inv and inv.get("rows"):
        ws = sheet("④출자 이력(내부·수동)", "출자 이력 – 내부 출자내역 기준(수동)", f"기준 해외자회사 출자내역 ({inv.get('updated', '')} 반영) · 단위: 현지통화, KRW 백만원, % · 회차별 송금일·현지통화 금액은 공시에 없음",
                   [26, 22, 12, 13, 12, 7, 20, 13, 11])
        for c, h in enumerate(["법인", "구분", "취득방법", "해외직접투자\n신고일", "송금일", "통화", "현지통화 금액", "원화 취득원가", "투자 후\n지분율"]):
            ws.write(3, 1 + c, h, hdr)
        ws.set_row(3, 30)
        r = 4
        for e in E:
            for x in e["hist"]:
                ws.write(r, 1, e["label"], lab)
                ws.write(r, 2, x["type"] + (" (사후보고)" if x.get("post_report") else ""), lab)
                for c, v in enumerate([x.get("method"), x.get("report_date"), x.get("remit_date"), x.get("ccy")]):
                    put(ws, r, 3 + c, v, txt)
                put(ws, r, 7, x.get("local_amt") or None, qty)
                put(ws, r, 8, x.get("krw_mn") or None, num)
                put(ws, r, 9, x.get("stake_after"), fm(border=1, num_format="0.00%", align="center"))
                r += 1
            if e["hist"]:
                cum = sum(x.get("krw_mn") or 0 for x in e["hist"])
                bk = (e["inv"] or {}).get("end_book")
                ws.merge_range(r, 1, r, 7, f"{e['label']} 누적 원화 취득원가" + (f" · DART 기말 장부가 {fmt(bk)} (차이 {fmt(bk - cum)})" if bk else ""), labb)
                put(ws, r, 8, cum, numb)
                ws.write_blank(r, 9, None, numb)
                r += 1

    # ⑤ 출처
    ws = sheet("⑤출처", "출처와 자동 갱신", "공시 자료만 사용 · 출자 이력만 내부 출자내역(수동)", [26, 70])
    rows = [(f"{p['label']} 실적·재무상태", p["report"]["name"], p["report"]["url"], ymd(p["report"]["rcept_dt"])) for p in D["periods"]]
    if ir:
        rows.append(("지분율·출자 현황", f"{ir['report']['name']} 타법인 출자현황", ir["report"]["url"], ymd(ir["report"]["rcept_dt"])))
    for i, (a, b, u, d) in enumerate(rows):
        ws.write(3 + i, 1, a, labb)
        ws.write_url(3 + i, 2, u, link, f"{b} (접수 {d})")
    r = 4 + len(rows)
    nx = M["next"]
    for i, s in enumerate([
        "실적 = 연결재무제표 주석 「종속기업 요약재무정보」(자산·부채·영업수익·당기순이익, 법인 단위·연결 조정 전).",
        "자동 갱신: GitHub Actions가 평일 07:30·16:30(KST) DART를 확인해 새 정기보고서가 있으면 화면과 이 파일을 다시 만든다(LLM 미사용).",
        f"다음 공시: {nx['item'].replace(' 해외자회사·타법인출자현황 반영', '')} (제출 기한 {nx['due']})" if nx else "다음 정기보고서 공시 후 자동 갱신.",
        "포함하지 않는 것: 2026 사업계획·계획 대비 진척률, 현지 결산(사업부문별 손익), 비공시 특이사항, 주재원 인원.",
    ]):
        ws.write(r + i, 1, s, note)
    wb.close()
    print("saved", XLSX.name)


# ---------- 1장 보고서 (Word, A4 가로) ----------
def build_docx(M):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    L, F, E, D = M["L"], M["F"], M["E"], M["D"]

    def shade(cell, hex_):
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_)
        cell._tc.get_or_add_tcPr().append(s)

    def run(p, text, size=8.5, bold=False, color=None):
        r = p.add_run(text)
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if color:
            r.font.color.rgb = rgb(color)
        return r

    def para(container, text="", size=8.5, bold=False, color=None, after=1):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
        if text:
            run(p, text, size, bold, color)
        return p

    def table(heads, rows, widths, left=1):
        t = doc.add_table(rows=1 + len(rows), cols=len(heads))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; t.autofit = False
        for i, h in enumerate(heads):
            c = t.rows[0].cells[i]; c.width = Cm(widths[i]); shade(c, NAVY.lstrip("#"))
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(c.paragraphs[0], h, 8, True, "#FFFFFF")
        for r, vals in enumerate(rows, 1):
            total = vals and str(vals[0]).startswith("4사")
            for i, v in enumerate(vals):
                c = t.rows[r].cells[i]; c.width = Cm(widths[i])
                if total:
                    shade(c, "F2F2F2")
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i < left else WD_ALIGN_PARAGRAPH.RIGHT
                s = str(v)
                run(c.paragraphs[0], s, 8.5, i == 0 or total, "#C00000" if s.startswith("(") or s.startswith("-") else None)
        return t

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin, sec.bottom_margin = Cm(1.1), Cm(0.9)

    p = para(doc, after=0)
    run(p, f"해외자회사 공시 실적 요약 ({L['label']})", 15, True, NAVY)
    run(p, f"    기준 DART {L['report']['name']}{' · ' + F['report']['name'] if F else ''} · 단위 KRW 백만원 · 미래전략실 글로벌네트워크팀", 8.5, False, GREY)
    a, b = summary_lines(M)
    p = para(doc, after=1)
    run(p, "■ 4사 합산  ", 10, True, ORANGE)
    run(p, a, 9.5, True)
    para(doc, b, 9, after=5)

    heads = ["법인 (지분율)", "주력 사업"] + ([f"{F['label']} 영업수익", f"{F['label']} 순이익", f"{F['label']} ROE"] if F else []) + \
            [f"{L['label']} 영업수익", f"{L['label']} 순이익", f"{L['label']} ROE*", "자본총계", "장부가액", "순자산/장부가"]
    rows = []
    for e in E:
        f, l, x = e["f"] or {}, e["l"] or {}, e["inv"] or {}
        nav = l["eq"] * e["stake"] if l.get("eq") is not None else None
        rows.append([f"{e['label']} {e['stake'] * 100:.{2 if 0.999 < e['stake'] < 1 else 0}f}%", e["biz"]]
                    + ([fmt(f.get("rev")), fmt(f.get("ni")), pct(f.get("roe"))] if F else [])
                    + [fmt(l.get("rev")), fmt(l.get("ni")), pct(l.get("roe")), fmt(l.get("eq")), fmt(x.get("end_book")),
                       pct(nav / x["end_book"], 0) if nav is not None and x.get("end_book") else "–"])
    tot, roe = M["tot"], M["roe"]
    sb = vsum(book_of(e) for e in E)
    nb = M["both"](nav_of, book_of)  # 순자산/장부가 합계는 둘 다 있는 법인끼리
    rows.append(["4사 합산 (100%)", ""] + ([fmt(tot("f", "rev")), fmt(tot("f", "ni")), pct(roe("f", F))] if F else [])
                + [fmt(tot("l", "rev")), fmt(tot("l", "ni")), pct(roe("l", L)), fmt(tot("l", "eq")), fmt(sb), pct(nb[0] / nb[1], 0) if nb and nb[1] else "–"])
    widths = [4.2, 4.0] + ([2.3, 2.1, 1.8] if F else []) + [2.3, 2.1, 1.9, 2.1, 2.1, 2.0]
    table(heads, rows, widths, left=2)
    para(doc, f"* ROE = 누적 순이익 × (12 ÷ {L['months']}) ÷ 기말 자본(자산 − 부채). 순자산/장부가 = 자본총계 × 지분율 ÷ 당사 장부가액(타법인 출자현황). "
              "요약재무정보는 법인 단위·연결 조정 전 금액.", 7.5, color=GREY, after=6)

    # 재무상태 변동
    per = ([F] if F else []) + [L]
    p = para(doc, after=2)
    run(p, "■ 재무상태  ", 10, True, ORANGE)
    run(p, " · ".join(f"{pp['label']} 말" for pp in per), 8, False, GREY)
    heads = ["법인"] + [f"{pp['label']} 말 {m}" for pp in per for m in ("자산", "부채", "자본")] + (["자본 증감"] if F else [])
    rows = []
    for e in E:
        vals = [fmt((e[sd] or {}).get(m)) for sd in (["f"] if F else []) + ["l"] for m in ("assets", "liab", "eq")]
        d = eq_of(e, "l") - eq_of(e, "f") if F and eq_of(e, "l") is not None and eq_of(e, "f") is not None else None
        rows.append([e["label"]] + vals + ([("+" if d and d > 0 else "") + fmt(d)] if F else []))
    dd = M["both"](lambda e: eq_of(e, "l"), lambda e: eq_of(e, "f")) if F else None
    dT = dd[0] - dd[1] if dd else None
    rows.append(["4사 합산"] + [fmt(tot(sd, m)) for sd in (["f"] if F else []) + ["l"] for m in ("assets", "liab", "eq")] + ([("+" if dT and dT > 0 else "") + fmt(dT)] if F else []))
    table(heads, rows, [4.6] + [3.4] * (3 * len(per)) + ([1.9] if F else []))
    para(doc, "", after=4)

    # 출자 현황: DART + 내부 출자 이력 요약
    ir = D.get("invest")
    p = para(doc, after=2)
    run(p, "■ 출자 현황  ", 10, True, ORANGE)
    run(p, f"DART {ir['report']['name']} 타법인 출자현황" if ir else "타법인 출자현황 없음", 8, False, GREY)
    run(p, " · 회차별 송금은 내부 출자내역 기준(수동)", 8, False, GREY)
    heads = ["법인", "최초 취득일", "최초 취득금액", "기초 장부가", "증가(감소)", "기말 장부가", "기말 지분율", "출자 회차", "최근 송금", "누적 원화 취득원가"]
    rows = []
    for e in E:
        x, hs = e["inv"] or {}, [h for h in e["hist"] if h.get("krw_mn")]
        last = max((h["remit_date"] for h in hs if h.get("remit_date")), default=None)
        rows.append([e["label"], (x.get("first_date") or "–").replace(".", "-"), fmt(x.get("first_amount")), fmt(x.get("begin_book")),
                     fmt(x.get("chg_amount")), fmt(x.get("end_book")), f"{x['end_stake']:.2f}%" if x.get("end_stake") is not None else "–",
                     f"{len(hs)}회" if hs else "–", last or "–", fmt(sum(h["krw_mn"] for h in hs)) if hs else "–"])
    rows.append(["4사 합계", "", "", fmt(vsum((e["inv"] or {}).get("begin_book") for e in E)), "", fmt(sb), "", "", "",
                 fmt(vsum(h.get("krw_mn") for e in E for h in e["hist"]))])
    table(heads, rows, [4.6, 2.4, 2.6, 2.6, 2.4, 2.6, 2.2, 1.9, 2.4, 3.2])

    p = para(doc, after=1); p.paragraph_format.space_before = Pt(8)
    run(p, "■ 다음 공시  ", 10, True, ORANGE)
    nx = M["next"]
    run(p, (f"{nx['item'].replace(' 해외자회사·타법인출자현황 반영', '')} 제출 기한 {nx['due']}. 공시되면 평일 자동 확인(07:30·16:30) 때 화면과 이 보고서가 새 기간으로 바뀝니다."
             if nx else "다음 정기보고서가 공시되면 자동으로 바뀝니다."), 9)
    p = para(doc, after=0)
    run(p, "■ 출처  ", 10, True, ORANGE)
    run(p, " · ".join(f"{pp['report']['name']} (접수 {ymd(pp['report']['rcept_dt'])})" for pp in D["periods"])
        + ". 공시 자료만 사용하며 사업계획·계획 대비, 현지 결산, 비공시 특이사항은 넣지 않았습니다.", 8, color=GREY)
    doc.save(str(DOCX))
    print("saved", DOCX.name)


def main(*flags):
    if "--check" in flags:
        print(f"rebuild={'true' if needs_rebuild() else 'false'}")
        return
    if not (load("dart_subs.json") or {}).get("periods"):
        print("dart_subs.json에 공시 기간이 없어 다운로드를 만들지 않음")
        return
    DL.mkdir(parents=True, exist_ok=True)
    M = model()
    build_xlsx(M)
    build_docx(M)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:  # 기준을 남기지 않아 LibreOffice가 있는 Actions가 다음에 PDF까지 다시 만든다
        print("LibreOffice가 없어 PDF는 건너뜀 (basis.json 미기록)")
        return
    PDF.unlink(missing_ok=True)  # 변환이 실패하면 예전 기간 PDF가 남지 않게 먼저 지운다
    r = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(DL), str(DOCX)], check=False,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    if r.returncode or not PDF.exists():
        sys.exit("PDF 변환 실패 (basis.json 미기록, 다음 실행에서 다시 만듦)")
    print("saved", PDF.name)
    BASIS.write_text(json.dumps(basis(), ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main(*sys.argv[1:])
